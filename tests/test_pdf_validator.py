import pytest

from app.services.pdf.validator import PdfValidationError, validate_pdf_upload


def test_valid_pdf_passes():
    validate_pdf_upload(file_size=1000, mime_type="application/pdf", header_bytes=b"%PDF-1.7")


def test_rejects_empty_file():
    with pytest.raises(PdfValidationError):
        validate_pdf_upload(file_size=0, mime_type="application/pdf", header_bytes=b"%PDF-1.7")


def test_rejects_oversized_file():
    from app.core.config import settings

    with pytest.raises(PdfValidationError):
        validate_pdf_upload(
            file_size=settings.pdf_max_size_bytes + 1, mime_type="application/pdf", header_bytes=b"%PDF-1.7"
        )


def test_rejects_wrong_mime_type():
    with pytest.raises(PdfValidationError):
        validate_pdf_upload(file_size=1000, mime_type="image/png", header_bytes=b"%PDF-1.7")


def test_rejects_missing_pdf_signature():
    with pytest.raises(PdfValidationError):
        validate_pdf_upload(file_size=1000, mime_type="application/pdf", header_bytes=b"NOT_A_PDF")


def test_valid_docx_passes():
    from app.services.pdf.validator import DOCX_MIME_TYPE, validate_docx_upload

    validate_docx_upload(file_size=1000, mime_type=DOCX_MIME_TYPE, header_bytes=b"PK\x03\x04\x14")


def test_docx_rejects_wrong_signature_and_mime():
    from app.services.pdf.validator import DOCX_MIME_TYPE, validate_docx_upload

    with pytest.raises(PdfValidationError):
        validate_docx_upload(file_size=1000, mime_type=DOCX_MIME_TYPE, header_bytes=b"%PDF-")
    with pytest.raises(PdfValidationError):
        validate_docx_upload(file_size=1000, mime_type="image/png", header_bytes=b"PK\x03\x04")


def test_extract_text_from_docx(tmp_path):
    import docx

    from app.services.pdf.extractor import extract_text

    path = tmp_path / "sample.docx"
    d = docx.Document()
    d.add_paragraph("Photosynthesis converts light into chemical energy in plants.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Chlorophyll"
    table.rows[0].cells[1].text = "green pigment"
    d.save(str(path))

    result = extract_text(path)
    assert "Photosynthesis" in result.text
    assert "Chlorophyll | green pigment" in result.text
    assert result.used_ocr is False
