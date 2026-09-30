import re
from pathlib import Path

import fitz  # PyMuPDF

from app.core.config import settings
from app.core.logging import get_logger

logger = get_logger(__name__)

MIN_CHARS_PER_PAGE_FOR_TEXT_LAYER = 20


class PdfExtractionError(Exception):
    pass


class PdfExtractionResult:
    def __init__(self, text: str, page_count: int, used_ocr: bool):
        self.text = text
        self.page_count = page_count
        self.used_ocr = used_ocr


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    return text.strip()


def _ocr_page(page: "fitz.Page") -> str:
    try:
        import pytesseract
        from PIL import Image
        import io

        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

        pix = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(image)
    except Exception as exc:  # OCR is best-effort; never crash the pipeline
        logger.warning("ocr_page_failed", error=str(exc))
        return ""


def _paragraph_text_with_bold_marks(paragraph) -> str:
    """Render a paragraph's text, wrapping bold runs in **markers**.

    Many exam DOCX files mark the correct answer option by bolding it
    instead of writing a separate answer key. Plain `paragraph.text`
    throws that formatting away, so the AI has no way to know which
    option is correct. Keeping a **bold** marker lets the bank-conversion
    prompt pick it up as a "highlighted option" mark.
    """
    # Merge consecutive runs with the same bold state first, so adjacent
    # bold runs (Word often splits "D) " and "600 ta" into separate runs)
    # become one "**D) 600 ta**" span instead of "**D) ****600 ta**".
    merged: list[tuple[bool, str]] = []
    for run in paragraph.runs:
        is_bold = bool(run.bold) if run.text.strip() else (merged[-1][0] if merged else False)
        if merged and merged[-1][0] == is_bold:
            merged[-1] = (is_bold, merged[-1][1] + run.text)
        else:
            merged.append((is_bold, run.text))

    parts = [f"**{text}**" if is_bold and text.strip() else text for is_bold, text in merged]
    return "".join(parts)


def _extract_docx(docx_path: Path) -> PdfExtractionResult:
    try:
        import docx

        document = docx.Document(str(docx_path))
    except Exception as exc:
        raise PdfExtractionError(f"Could not open DOCX: {exc}") from exc

    parts = [_paragraph_text_with_bold_marks(p) for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    # Join with a blank line between paragraphs so downstream paragraph-based
    # chunking (split_into_chunks) and line-based bank detection (analyzer.py)
    # both see real paragraph/line boundaries instead of one giant blob.
    full_text = _clean_text("\n\n".join(parts))
    if not full_text:
        raise PdfExtractionError("No extractable text found in DOCX")

    # DOCX has no fixed pages; approximate for stats only.
    page_count = max(1, len(full_text) // 3000)
    return PdfExtractionResult(text=full_text, page_count=page_count, used_ocr=False)


def extract_text(pdf_path: Path) -> PdfExtractionResult:
    if Path(pdf_path).suffix.lower() == ".docx":
        return _extract_docx(Path(pdf_path))

    try:
        doc = fitz.open(pdf_path)
    except Exception as exc:
        raise PdfExtractionError(f"Could not open PDF: {exc}") from exc

    if doc.is_encrypted:
        try:
            doc.authenticate("")
        except Exception:
            raise PdfExtractionError("PDF is password-protected")

    page_texts: list[str] = []
    used_ocr = False

    for page in doc:
        text = page.get_text("text")
        if len(text.strip()) < MIN_CHARS_PER_PAGE_FOR_TEXT_LAYER and settings.ocr_enabled:
            ocr_text = _ocr_page(page)
            if len(ocr_text.strip()) > len(text.strip()):
                text = ocr_text
                used_ocr = True
        page_texts.append(text)

    page_count = doc.page_count
    doc.close()

    full_text = _clean_text("\n\n".join(page_texts))
    if not full_text:
        raise PdfExtractionError("No extractable text found in PDF, even after OCR")

    return PdfExtractionResult(text=full_text, page_count=page_count, used_ocr=used_ocr)
